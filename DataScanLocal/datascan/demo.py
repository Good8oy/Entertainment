"""Synthetic, explicitly labelled examples; never presented as real research data."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def create_demo(directory: Path) -> list[str]:
    directory.mkdir(parents=True, exist_ok=True)
    values = [12.3847, 18.9263, 15.4718, 24.6392, 20.8174, 31.2569, 27.5931, 35.7486]
    first = Workbook()
    ws = first.active
    ws.title = "实验记录_演示"
    ws.append(["Time (h)", "Signal A", "Replicate 1", "Replicate 2", "周期示例"])
    cycle = [2.184]
    for i in range(12):
        cycle.append(round(cycle[-1] + (0.371, 0.823, 0.146)[i % 3], 4))
    for i in range(13):
        value = values[i] if i < len(values) else None
        ws.append([i * 2, value, value, round(value + 0.25, 4) if value else None, cycle[i]])
    first.create_sheet("矩形源").append(["演示合成数据 / 非真实实验"])
    for row in ([3.172, 8.463, 4.291], [9.537, 2.684, 7.815]):
        first["矩形源"].append(row)
    second = Workbook()
    ws = second.active
    ws.title = "跨文件对照_演示"
    ws.append(["Index", "完全复用", "比例示例", "整数改写", "重排示例", "稀疏复用"])
    order = [3, 1, 6, 0, 7, 2, 5, 4]
    for i in range(16):
        v = values[i] if i < 8 else None
        ws.append([i + 1, v, round(v * 2.5, 6) if v else None,
                   round(v + (i + 1) * 10, 4) if v else None,
                   values[order[i]] if i < 8 else None,
                   values[i // 2] if i % 2 == 0 else 90.027 + i * 0.3127])
    ws = second.create_sheet("矩形目标")
    ws.append(["独立构造的命中示例"])
    for row in ([3.172, 8.463, 4.291], [9.537, 2.684, 7.815]):
        ws.append(row)
    ws = second.create_sheet("末位分布_演示")
    ws.append(["合成偏态末位样本"])
    from random import Random
    random = Random(42)
    for i in range(60):
        ws.append([round(random.randrange(1, 10000) / 100 + 0.001, 3)])
    paths = []
    for filename, book in [("演示_A.xlsx", first), ("演示_B.xlsx", second)]:
        for sheet in book:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="315ED5")
            for col in "ABCDEF":
                sheet.column_dimensions[col].width = 24
        path = directory / filename
        book.save(path)
        book.close()
        paths.append(str(path))
    return paths
