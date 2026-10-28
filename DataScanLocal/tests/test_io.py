import csv
import hashlib
import json

import openpyxl
import pytest

from datascan.engine import scan
from datascan.loader import load_file
from datascan.model import Evidence, Region, ScanResult, Settings, Sheet
from datascan.report import export_report


def test_xlsx_multiple_hidden_sheets_and_formula_exclusion(tmp_path):
    path = tmp_path / "book.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["Signal", "Calculated"])
    book.active.append([12.31, "=A2*2"])
    hidden = book.create_sheet("hidden")
    hidden.append([23.41])
    hidden.sheet_state = "hidden"
    book.save(path)
    sheets = load_file(str(path))
    assert len(sheets) == 2
    assert sheets[0].rows[1][1] == "=A2*2"
    assert "公式" in sheets[0].warnings[0]
    assert "隐藏" in sheets[1].warnings[0]
    assert sheets[0].source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.mark.parametrize("suffix,delimiter,encoding", [(".csv", ",", "utf-8-sig"), (".tsv", "\t", "utf-8"),
                                                       (".csv", ";", "gb18030"), (".csv", ",", "utf-16")])
def test_delimited_encodings_and_quoted_newline(tmp_path, suffix, delimiter, encoding):
    path = tmp_path / ("表格" + suffix)
    with path.open("w", encoding=encoding, newline="") as file:
        writer = csv.writer(file, delimiter=delimiter)
        writer.writerow(["样本", "数值"])
        writer.writerow(["第一行\n第二行", "12.3847"])
    sheet = load_file(str(path))[0]
    assert sheet.rows[1] == ["第一行\n第二行", "12.3847"]


def test_duplicate_basename_has_distinct_sheet_keys(tmp_path):
    a, b = tmp_path / "A", tmp_path / "B"
    a.mkdir()
    b.mkdir()
    for folder in (a, b):
        (folder / "data.csv").write_text("X\n1.23\n", encoding="utf-8")
    assert load_file(str(a / "data.csv"))[0].key != load_file(str(b / "data.csv"))[0].key


def test_even_length_gb18030_not_misread_as_utf16(tmp_path):
    path = tmp_path / "legacy.csv"
    raw = "样本,数值\n甲,1.23\n".encode("gb18030")
    if len(raw) % 2:
        raw += b"\n"
    path.write_bytes(raw)
    assert load_file(str(path))[0].rows[0] == ["样本", "数值"]


def test_invalid_encrypted_or_corrupt_excel_is_actionable(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"invalid")
    with pytest.raises(ValueError, match="加密或损坏"):
        load_file(str(path))


def test_xls_saved_numeric_values(tmp_path):
    import xlwt
    path = tmp_path / "legacy.xls"
    book = xlwt.Workbook()
    ws = book.add_sheet("原始数据")
    ws.write(0, 0, "Signal")
    ws.write(1, 0, 12.3847)
    ws.write(1, 1, True)
    book.save(str(path))
    sheet = load_file(str(path))[0]
    assert sheet.rows[1][0] == 12.3847
    assert isinstance(sheet.rows[1][1], str)
    assert "缓存" in sheet.warnings[0]


def result_fixture():
    sheet = Sheet("a", "<file>.csv", "<script>alert(1)</script>", [[1.23]])
    region = Region("a", ((1, 1),), (1.23,))
    evidence = Evidence("DS-0001", "duplicate", "<img src=x>", "高风险", "test", "=SUM(A1:A2)", [region],
                        note="=HYPERLINK(\"bad\")")
    return ScanResult(Settings(), [sheet], [evidence], ["<note>"], .01)


def test_exports_are_escaped_and_preserve_reviews(tmp_path):
    result = result_fixture()
    result.evidence[0].review = "正常共用／已解释"
    for suffix in ("html", "json", "csv"):
        export_report(result, str(tmp_path / ("report." + suffix)))
    html = (tmp_path / "report.html").read_text(encoding="utf-8")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    payload = json.loads((tmp_path / "report.json").read_text(encoding="utf-8"))
    assert payload["evidence"][0]["review"] == "正常共用／已解释"
    assert payload["evidence"][0]["regions"][0]["cells"] == [[1, 1]]
    assert payload["scope"][0]["path"] == "<file>.csv"
    text = (tmp_path / "report.csv").read_text(encoding="utf-8-sig")
    assert "'=SUM" in text
    assert "'=HYPERLINK" in text


def test_empty_report_still_exports_scope_and_limits(tmp_path):
    result = scan([Sheet("a", "data.csv", "Data", [["value"], [3.14]])])
    path = tmp_path / "empty.html"
    export_report(result, str(path))
    text = path.read_text(encoding="utf-8")
    assert "未检出" in text
    assert "核验范围" in text
    assert "方法与边界" in text


def test_csv_cell_limit_is_explicit(tmp_path, monkeypatch):
    import datascan.loader as loader
    monkeypatch.setattr(loader, "MAX_CELLS", 3)
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ValueError, match="超过"):
        load_file(str(path))
